"""Archivos mínimos generados en memoria para las pruebas de extracción."""

from io import BytesIO

from docx import Document
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def pdf_bytes(texts, *, encrypted=False):
    writer = PdfWriter()
    for text in texts:
        page = writer.add_blank_page(width=612, height=792)
        if text:
            font = DictionaryObject({NameObject("/Type"): NameObject("/Font"),
                                     NameObject("/Subtype"): NameObject("/Type1"),
                                     NameObject("/BaseFont"): NameObject("/Helvetica")})
            page[NameObject("/Resources")] = DictionaryObject({
                NameObject("/Font"): DictionaryObject({NameObject("/F1"): font})})
            stream = DecodedStreamObject()
            stream.set_data(f"BT /F1 12 Tf 40 700 Td ({text}) Tj ET".encode("ascii"))
            page[NameObject("/Contents")] = stream
    if encrypted:
        writer.encrypt("test-password")
    output = BytesIO()
    writer.write(output)
    return output.getvalue()


def docx_bytes(text="La escuderia Cometa gano la carrera ficticia del Circuito Lunar."):
    document = Document()
    document.add_paragraph(text)
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Piloto ficticio"
    table.cell(0, 1).text = "Elena Rayo"
    output = BytesIO()
    document.save(output)
    return output.getvalue()
